"""產生測試用的假資料 Excel（不含任何公司資料）。

產出到 tests/fixtures/（已被 .gitignore 排除，不會進 repo）：
  plain.xlsx       約 1 千筆，沒有密碼
  mid_plain.xlsx   約 5 萬筆，沒有密碼
  big_plain.xlsx   約 18 萬筆，沒有密碼（--big 才產生）
  agile.xlsx       plain.xlsx 加上 Excel 2010+ 的開啟密碼（AES-256 + SHA-512）
  big_agile.xlsx   big_plain.xlsx 加密（--big 才產生）
  standard.xlsx    plain.xlsx 加上 Excel 2007 的開啟密碼（AES-128 + SHA-1）
  month.xlsx       每月加入用的新月份資料（115/01 新增、114/12 取代）
  month_agile.xlsx month.xlsx 加密
  class.xlsx       假的商品類別代碼表；class_agile.xlsx 是加密版
  mid_standard.xlsx

需要：pip install openpyxl cryptography
用法：python tests/make_fixtures.py [--big]
"""
import base64, hashlib, hmac, os, random, struct, sys
import openpyxl
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes

HERE = os.path.dirname(os.path.abspath(__file__))
FX = os.path.join(HERE, 'fixtures')
AGILE_PW = '測試pw123'      # agile.xlsx 的密碼（含中文，順便測 UTF-16）
STANDARD_PW = 'Abc!2007'   # standard.xlsx 的密碼


# 假的商品：(類別代碼, 商品代號樣式, 尺寸, 售價)。代號寫法涵蓋「-」「/」和直接接在數字後面（AR9006XXL、6872XL-XXL）
PRODUCTS = [
    ('54', '5401-{}', ['M', 'L', 'XL'], 150),
    ('60', '6001/{}', ['M-L', 'XL-XXL'], 120),
    ('6-', '6872{}', ['M-L', 'XL-XXL'], 200),
    ('25', 'AR9006{}', ['L', 'XL', 'XXL'], 300),
    ('79', '7901-{}', ['90', '100', '110'], 100),
    ('92', '9200', [''], 50),          # 沒有尺寸（像運費）
    ('94', '9401-{}', ['F'], 80),      # 分類表裡沒有名稱的代碼
]
# 客戶：(代號, 名稱, 業務員)。9 開頭 → 其他地區；丁行 114 年起沒再買（可能流失）
CUSTS = [('1001', '甲商行', '01'), ('1002', '己百貨', '01'), ('3002', '乙公司', '02'), ('5003', '丙店', '03'), ('7004', '丁行', '02'), ('9005', '庚批發', '03')]
COLORS = ['01', '02', '05']


def make_xlsx(path, per_month):
    random.seed(1)
    wb = openpyxl.Workbook()
    wb.active.title = '樞紐'
    wb.active['A1'] = '這張工作表應該被忽略'
    s = wb.create_sheet('原始進貨表')
    s.append(['出貨日期', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '品項總額', '客戶代號', '客戶名', '業務員'])
    n = 0
    for y in (112, 113, 114):
        cost_ratio = 0.6 if y < 114 else 0.65   # 114 年毛利率變差
        custs = [c for c in CUSTS if not (y == 114 and c[1] == '丁行')]
        for m in range(1, 13):
            for k in range(per_month):
                c = random.choice(custs)
                cat, pat, sizes, price = random.choices(PRODUCTS, weights=[5, 4, 3, 3, 2, 1, 1])[0]
                size = random.choices(sizes, weights=list(range(len(sizes), 0, -1)))[0]   # 第一個尺寸最好賣
                color = random.choices(COLORS, weights=[3, 2, 1])[0]                       # 01 最好賣
                sku = pat.format(size) if '{}' in pat else pat
                q = random.randint(1, 20)
                unit, unitn = ('打', 2) if cat == '79' and k % 7 == 0 else ('件', 1)
                s.append([f'{y}{m:02d}{k:04d}', cat, sku, color, unit, round(price * cost_ratio, 2), price, q, unitn, q * price, c[0], c[1], c[2]])
                n += 1
    r = wb.create_sheet('原始退貨表')
    r.append(['退貨單號', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '退貨金額', '客戶代號', '客戶名', '業務員'])
    for y in (113, 114):
        for m in range(1, 13, 3):
            r.append([f'{y}{m:02d}900', '54', '5401-M', '01', '件', 90, 150, 2, 1, 300 if y == 113 else 3000, '1001', '甲商行', '01'])
    wb.save(path)
    print(f'{os.path.basename(path)}: {n} 筆進貨、8 筆退貨')


def make_class_xlsx(path):
    """假的商品類別代碼表（和使用者的格式一樣：商品類別(代碼)、商品類別、大類）。名稱都是假的。"""
    wb = openpyxl.Workbook()
    s = wb.active; s.title = '商品類別代碼'
    s.append(['商品類別(代碼)', '商品類別', '大類'])
    for row in [('54', '測試男褲', '男'), ('60', '測試女褲', '女'), ('6-', '測試女無縫', '女'), ('25', '測試男衣', '男'),
                ('79', '測試童褲', '童'), ('92', '測試運費', '配件'), (3, '測試手帕', '配件'), ('94', None, None), ('(空白)', None, None)]:
        s.append(list(row))
    wb.create_sheet('工作表2')
    wb.save(path)
    print(f'{os.path.basename(path)}: 商品類別代碼表（假資料）')


def make_month_xlsx(path):
    """每月加入用：115 年 1 月（新月份）25 筆進貨、2 筆退貨，加上 114 年 12 月（取代原有的 30 筆）10 筆進貨。"""
    wb = openpyxl.Workbook()
    s = wb.active; s.title = '原始進貨表'
    s.append(['出貨日期', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '品項總額', '客戶代號', '客戶名', '業務員'])
    for k in range(10):
        s.append([f'11412{k:04d}', '類A', 'A100-1', 'R', '雙', 60, 100, 1, 1, 100, '1001', '甲商行', '王'])
    for k in range(25):
        s.append([f'11501{k:04d}', '類D', 'D400-1', 'B', '雙', 120, 200, 2, 1, 400, '8005', '戊新店', '李'])
    r = wb.create_sheet('原始退貨表')
    r.append(['退貨單號', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '退貨金額', '客戶代號', '客戶名', '業務員'])
    for k in range(2):
        r.append([f'11501{900 + k}', '類D', 'D400-1', 'B', '雙', 120, 200, 1, 1, 200, '8005', '戊新店', '李'])
    wb.save(path)
    print(f'{os.path.basename(path)}: 每月加入用，35 筆進貨、2 筆退貨')


# ---------- 最小的 CFB（OLE 複合文件）寫入器 ----------
def write_cfb(path, info, pkg):
    SS = 512; END = 0xFFFFFFFE; FREE = 0xFFFFFFFF; FATSECT = 0xFFFFFFFD; DIFSECT = 0xFFFFFFFC
    mini_data = info + b'\0' * ((64 - len(info) % 64) % 64); n_mini = len(mini_data) // 64
    mini_secs = (len(mini_data) + SS - 1) // SS
    pkg_secs = (len(pkg) + SS - 1) // SS
    body = 2 + mini_secs + pkg_secs
    nfat = 1; ndif = 0
    while True:
        ndif = max(0, -(-(nfat - 109) // 127))
        if nfat * 128 >= body + nfat + ndif: break
        nfat += 1
    fat = []
    dif_sids = list(range(ndif)); fat += [DIFSECT] * ndif
    fat_sids = list(range(ndif, ndif + nfat)); fat += [FATSECT] * nfat
    dir_sid = len(fat); fat.append(END)
    mfat_sid = len(fat); fat.append(END)
    ms_sid = len(fat)
    for i in range(mini_secs): fat.append(len(fat) + 1 if i < mini_secs - 1 else END)
    pk_sid = len(fat)
    for i in range(pkg_secs): fat.append(len(fat) + 1 if i < pkg_secs - 1 else END)
    fat += [FREE] * (nfat * 128 - len(fat))

    def entry(name, typ, start, size, child=FREE, right=FREE):
        n = name.encode('utf-16le') + b'\0\0'
        return n + b'\0' * (64 - len(n)) + struct.pack('<HBB', len(n), typ, 1) + struct.pack('<III', FREE, right, child) + b'\0' * 36 + struct.pack('<II', start, size) + b'\0' * 4
    d = entry('Root Entry', 5, ms_sid, len(mini_data), child=1) + entry('EncryptionInfo', 2, 0, len(info), right=2) + entry('EncryptedPackage', 2, pk_sid, len(pkg)) + b'\0' * 128
    mfat = [i + 1 if i < n_mini - 1 else END for i in range(n_mini)] + [FREE] * (128 - n_mini)
    head = bytes.fromhex('d0cf11e0a1b11ae1') + b'\0' * 16 + struct.pack('<HHHHH', 0x3E, 3, 0xFFFE, 9, 6) + b'\0' * 6
    head += struct.pack('<IIIIIIIII', 0, nfat, dir_sid, 0, 4096, mfat_sid, 1, dif_sids[0] if ndif else END, ndif)
    first = fat_sids[:109]
    head += b''.join(struct.pack('<I', x) for x in first) + b''.join(struct.pack('<I', FREE) for _ in range(109 - len(first)))
    difs = b''
    rest = fat_sids[109:]
    for k in range(ndif):
        chunk = rest[k * 127:(k + 1) * 127]
        nxt = dif_sids[k + 1] if k + 1 < ndif else END
        difs += b''.join(struct.pack('<I', x) for x in chunk) + b''.join(struct.pack('<I', FREE) for _ in range(127 - len(chunk))) + struct.pack('<I', nxt)
    pad = lambda b: b + b'\0' * ((SS - len(b) % SS) % SS)
    out = head + difs + b''.join(struct.pack('<I', x) for x in fat) + d + b''.join(struct.pack('<I', x) for x in mfat) + pad(mini_data) + pad(pkg)
    open(path, 'wb').write(out)


def aes(key, mode, data, enc=True):
    c = Cipher(algorithms.AES(key), mode)
    x = c.encryptor() if enc else c.decryptor()
    return x.update(data) + x.finalize()


def encrypt_standard(src, dst, pw):
    """Excel 2007 的「Standard」加密（[MS-OFFCRYPTO] 2.3.4.5）。"""
    plain = open(src, 'rb').read()
    salt = os.urandom(16)
    h = hashlib.sha1(salt + pw.encode('utf-16le')).digest()
    for i in range(50000): h = hashlib.sha1(struct.pack('<I', i) + h).digest()
    hf = hashlib.sha1(h + struct.pack('<I', 0)).digest()
    b1 = bytes(x ^ 0x36 for x in hf) + b'\x36' * 44; b2 = bytes(x ^ 0x5c for x in hf) + b'\x5c' * 44
    key = (hashlib.sha1(b1).digest() + hashlib.sha1(b2).digest())[:16]
    ver = os.urandom(16)
    encVer = aes(key, modes.ECB(), ver); encVH = aes(key, modes.ECB(), hashlib.sha1(ver).digest() + b'\0' * 12)
    csp = 'Microsoft Enhanced RSA and AES Cryptographic Provider\0'.encode('utf-16le')
    hdr = struct.pack('<IIIIIIII', 0x24, 0, 0x660E, 0x8004, 128, 0x18, 0, 0) + csp
    info = struct.pack('<HHI', 3, 2, 0x24) + struct.pack('<I', len(hdr)) + hdr + struct.pack('<I', 16) + salt + encVer + struct.pack('<I', 20) + encVH
    pkg = struct.pack('<Q', len(plain)) + aes(key, modes.ECB(), plain + b'\0' * ((16 - len(plain) % 16) % 16))
    write_cfb(dst, info, pkg)
    print(f'{os.path.basename(dst)}: Standard 加密，密碼 {pw}')


def encrypt_agile(src, dst, pw, spin=100000):
    """Excel 2010+ 的「Agile」加密（[MS-OFFCRYPTO] 2.3.4.10–2.3.4.15），AES-256-CBC + SHA-512。"""
    plain = open(src, 'rb').read()
    H = lambda b: hashlib.sha512(b).digest()
    fix = lambda b, n: b[:n] if len(b) >= n else b + b'\x36' * (n - len(b))
    secret = os.urandom(32); kd_salt = os.urandom(16); ek_salt = os.urandom(16)
    # 內容：每 4096 位元組一段，各自的 IV
    out = b''
    for i in range(0, len(plain), 4096):
        seg = plain[i:i + 4096]; seg += b'\0' * ((16 - len(seg) % 16) % 16)
        iv = fix(H(kd_salt + struct.pack('<I', i // 4096)), 16)
        out += aes(secret, modes.CBC(iv), seg)
    pkg = struct.pack('<Q', len(plain)) + out
    # 由密碼產生的金鑰
    h = H(ek_salt + pw.encode('utf-16le'))
    for i in range(spin): h = H(struct.pack('<I', i) + h)
    bk = lambda block: fix(H(h + block), 32)
    iv = fix(ek_salt, 16)
    vin = os.urandom(16)
    e_vin = aes(bk(bytes.fromhex('fea7d2763b4b9e79')), modes.CBC(iv), vin)
    e_vh = aes(bk(bytes.fromhex('d7aa0f6d3061344e')), modes.CBC(iv), H(vin))
    e_key = aes(bk(bytes.fromhex('146e0be7abacd0d6')), modes.CBC(iv), secret)
    # 資料完整性 HMAC（軟體不檢查，但 Excel 需要）
    hk = os.urandom(64)
    iv1 = fix(H(kd_salt + bytes.fromhex('5fb2ad010cb9e1f6')), 16)
    iv2 = fix(H(kd_salt + bytes.fromhex('a0677f02b22c8433')), 16)
    e_hk = aes(secret, modes.CBC(iv1), hk)
    e_hv = aes(secret, modes.CBC(iv2), hmac.new(hk, pkg, hashlib.sha512).digest())
    b = lambda x: base64.b64encode(x).decode()
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
           '<encryption xmlns="http://schemas.microsoft.com/office/2006/encryption" xmlns:p="http://schemas.microsoft.com/office/2006/keyEncryptor/password">'
           f'<keyData saltSize="16" blockSize="16" keyBits="256" hashSize="64" cipherAlgorithm="AES" cipherChaining="ChainingModeCBC" hashAlgorithm="SHA512" saltValue="{b(kd_salt)}"/>'
           f'<dataIntegrity encryptedHmacKey="{b(e_hk)}" encryptedHmacValue="{b(e_hv)}"/>'
           '<keyEncryptors><keyEncryptor uri="http://schemas.microsoft.com/office/2006/keyEncryptor/password">'
           f'<p:encryptedKey spinCount="{spin}" saltSize="16" blockSize="16" keyBits="256" hashSize="64" cipherAlgorithm="AES" cipherChaining="ChainingModeCBC" hashAlgorithm="SHA512" saltValue="{b(ek_salt)}" '
           f'encryptedVerifierHashInput="{b(e_vin)}" encryptedVerifierHashValue="{b(e_vh)}" encryptedKeyValue="{b(e_key)}"/>'
           '</keyEncryptor></keyEncryptors></encryption>')
    info = struct.pack('<HHI', 4, 4, 0x40) + xml.encode('utf-8')
    write_cfb(dst, info, pkg)
    print(f'{os.path.basename(dst)}: Agile 加密，密碼 {pw}')


# ---------- 系統原始資料（JMSSAL 出貨、JMSSBK 退貨）：英文欄位，W／V 欄是實際日期 ----------
SAL_FIELDS = [('DELIVNO', 'C', 10, 0), ('DELIVORD', 'C', 2, 0), ('CODE', 'C', 2, 0), ('STYLE', 'C', 15, 0), ('COLO', 'C', 2, 0), ('UNIT_NM', 'C', 2, 0),
              ('TRAN_RATE', 'N', 4, 0), ('PRICE', 'N', 10, 1), ('OUTPRICE', 'N', 10, 1), ('PRC', 'N', 10, 1), ('OUTQTY', 'N', 8, 1), ('UNIT', 'C', 1, 0),
              ('SALEGET', 'N', 12, 1), ('RECEIVE', 'N', 12, 1), ('CUT', 'L', 1, 0), ('TOTBKPAY', 'N', 12, 1), ('PAY', 'L', 1, 0), ('BUY', 'L', 1, 0),
              ('ID', 'C', 6, 0), ('SNAME', 'C', 20, 0), ('AMOUNT', 'N', 12, 0), ('EMPNO', 'C', 3, 0), ('DELIVDATE', 'C', 7, 0)]
BK_FIELDS = [('DELIVNO', 'C', 10, 0), ('SERIES', 'C', 4, 0), ('DELIVORD', 'C', 2, 0), ('CODE', 'C', 2, 0), ('STYLE', 'C', 15, 0), ('COLO', 'C', 2, 0), ('UNIT_NM', 'C', 2, 0),
             ('TRAN_RATE', 'N', 4, 0), ('PRICE', 'N', 10, 1), ('OUTPRICE', 'N', 10, 1), ('OUTQTY', 'N', 8, 1), ('UNIT', 'C', 1, 0), ('SALEGET', 'N', 12, 1),
             ('RECEIVE', 'N', 12, 1), ('TOTBKPAY', 'N', 12, 1), ('PAY', 'L', 1, 0), ('HOLD', 'L', 1, 0), ('ID', 'C', 6, 0), ('SNAME', 'C', 20, 0),
             ('AMOUNT', 'N', 12, 0), ('EMPNO', 'C', 3, 0), ('DELIVDATE', 'C', 7, 0)]
RAW_CUSTS = [('2T06', '甲商行'), ('5X03', '乙公司'), ('3A11', '丙店'), ('7K02', '丁行')]


def raw_rows():
    """回傳 (出貨列, 退貨列)，每列是 dict。114/09～114/11；單號 11411… 的單有一部分實際在 10/30、10/31 出貨。"""
    random.seed(7)
    sal, bk = [], []
    no = 0
    for ym in (11409, 11410, 11411):
        for k in range(40):
            no += 1
            cid, name = random.choice(RAW_CUSTS)
            day = random.randint(1, 28)
            date = ym * 100 + day
            if ym == 11411 and k < 10: date = 11410 * 100 + 30 + (k % 2)   # 11 月的單號，10 月底就出貨
            delivno = f'{ym}{no:05d}'
            for line in range(1, random.randint(2, 3) + 1):
                cat, pat, sizes, price = random.choice(PRODUCTS[:5])
                sku = pat.format(random.choice(sizes))
                dozen = random.random() < 0.5
                q = random.choice([0.5, 1, 2]) if dozen else random.randint(1, 10)
                amt = round(q * price * (12 if dozen else 1), 1)
                sal.append(dict(DELIVNO=delivno, DELIVORD=f'{line:02d}', CODE=cat, STYLE=sku, COLO=random.choice(COLORS[:2]), UNIT_NM='打' if dozen else '件',
                                TRAN_RATE=12 if dozen else 1, PRICE=round(price * (12 if dozen else 1) * 0.7, 1), OUTPRICE=price * (12 if dozen else 1), PRC=0.0, OUTQTY=q,
                                UNIT='2' if dozen else '1', SALEGET=amt, RECEIVE=0.0, CUT=False, TOTBKPAY=0.0, PAY=False, BUY=False, ID=cid, SNAME=name,
                                AMOUNT=0, EMPNO='22', DELIVDATE=str(date)))
    for i in range(6):   # 退貨：9 月單號但 10 月 2 日才退回 2 筆
        ym = 11409 + (i // 2)
        date = (11410 * 100 + 2) if i < 2 else ym * 100 + 15
        cid, name = RAW_CUSTS[i % 4]
        bk.append(dict(DELIVNO=f'{ym}9{i:04d}', SERIES='', DELIVORD='01', CODE='54', STYLE='5401-M', COLO='01', UNIT_NM='件', TRAN_RATE=1, PRICE=105.0,
                       OUTPRICE=150.0, OUTQTY=2.0, UNIT='1', SALEGET=300.0, RECEIVE=0.0, TOTBKPAY=0.0, PAY=False, HOLD=True, ID=cid, SNAME=name,
                       AMOUNT=300, EMPNO='21', DELIVDATE=str(date)))
    return sal, bk


def write_dbf(path, fields, rows, deleted=()):
    """dBase III 檔，中文用 Big5（和系統匯出的 DBF 一樣）。deleted：要標成已刪除的列號。"""
    rec_len = 1 + sum(f[2] for f in fields)
    head_len = 32 + 32 * len(fields) + 1
    out = bytearray(struct.pack('<BBBBIHH20x', 0x03, 126, 10, 30, len(rows), head_len, rec_len))
    for name, typ, ln, dec in fields:
        out += name.encode('ascii').ljust(11, b'\0') + typ.encode() + b'\0' * 4 + bytes([ln, dec]) + b'\0' * 14
    out += b'\r'
    for i, r in enumerate(rows):
        out += b'*' if i in deleted else b' '
        for name, typ, ln, dec in fields:
            v = r[name]
            if typ == 'C': b = str(v).encode('big5').ljust(ln, b' ')[:ln]
            elif typ == 'L': b = b'T' if v else b'F'
            else: b = (f'{v:.{dec}f}' if dec else f'{int(v)}').rjust(ln).encode()
            out += b
    out += b'\x1a'
    open(path, 'wb').write(bytes(out))


def write_raw_xlsx(path, sheet, fields, rows):
    """像 WPS 把 DBF 另存成 Excel：工作表名稱是檔名、英文表頭、文字欄位存成文字。"""
    wb = openpyxl.Workbook(); s = wb.active; s.title = sheet
    s.append([f[0] for f in fields])
    for r in rows: s.append([r[f[0]] for f in fields])
    wb.save(path)


def make_raw(fx):
    sal, bk = raw_rows()
    write_dbf(os.path.join(fx, 'JMSSAL.DBF'), SAL_FIELDS, sal + [dict(sal[0], DELIVNO='1141199999')], deleted={len(sal)})
    write_dbf(os.path.join(fx, 'JMSSBK.DBF'), BK_FIELDS, bk)
    write_raw_xlsx(os.path.join(fx, 'JMSSAL.xlsx'), 'JMSSAL', SAL_FIELDS, sal)
    write_raw_xlsx(os.path.join(fx, 'JMSSBK.xlsx'), 'JMSSBK', BK_FIELDS, bk)
    print(f'JMSSAL／JMSSBK（.DBF 與 .xlsx）: {len(sal)} 筆出貨（DBF 另有 1 筆已刪除）、{len(bk)} 筆退貨')



def make_dated_master(path):
    """總表更新工具加上實際日期欄之後的總表：中文表頭，原始進貨表最後有「實際出貨日期」、原始退貨表有「實際退貨日期」。內容和 raw_rows() 一樣。"""
    sal, bk = raw_rows()
    wb = openpyxl.Workbook(); wb.active.title = '樞紐'
    s = wb.create_sheet('原始進貨表')
    s.append(['出貨日期', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '品項總額', '客戶代號', '客戶名', '業務員', '實際出貨日期'])
    for r in sal: s.append([r['DELIVNO'], r['CODE'], r['STYLE'], r['COLO'], r['UNIT_NM'], r['PRICE'], r['OUTPRICE'], r['OUTQTY'], int(r['UNIT']), r['SALEGET'], r['ID'], r['SNAME'], int(r['EMPNO']), r['DELIVDATE']])
    t = wb.create_sheet('原始退貨表')
    t.append(['退貨單號', '商品類別', '商品代號', '顏色代號', '單位', '單價', '售價', '數量', '單位(數字)', '退貨金額', '客戶代號', '客戶名', '業務員', '實際退貨日期'])
    for r in bk: t.append([r['DELIVNO'], r['CODE'], r['STYLE'], r['COLO'], r['UNIT_NM'], r['PRICE'], r['OUTPRICE'], r['OUTQTY'], int(r['UNIT']), r['SALEGET'], r['ID'], r['SNAME'], int(r['EMPNO']), r['DELIVDATE']])
    wb.save(path)
    print(f'{os.path.basename(path)}: 有實際日期欄的總表')

if __name__ == '__main__':
    os.makedirs(FX, exist_ok=True)
    p = lambda n: os.path.join(FX, n)
    make_xlsx(p('plain.xlsx'), 30)
    make_xlsx(p('mid_plain.xlsx'), 1500)
    make_month_xlsx(p('month.xlsx'))
    make_class_xlsx(p('class.xlsx'))
    make_raw(FX)
    make_dated_master(p('master_dated.xlsx'))
    encrypt_agile(p('class.xlsx'), p('class_agile.xlsx'), AGILE_PW)
    encrypt_agile(p('month.xlsx'), p('month_agile.xlsx'), AGILE_PW)
    encrypt_agile(p('plain.xlsx'), p('agile.xlsx'), AGILE_PW)
    encrypt_standard(p('plain.xlsx'), p('standard.xlsx'), STANDARD_PW)
    encrypt_standard(p('mid_plain.xlsx'), p('mid_standard.xlsx'), STANDARD_PW)
    if '--big' in sys.argv:
        make_xlsx(p('big_plain.xlsx'), 5000)
        encrypt_agile(p('big_plain.xlsx'), p('big_agile.xlsx'), AGILE_PW)
